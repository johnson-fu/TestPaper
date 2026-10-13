#!/usr/bin/env python3
"""依 JSON 設定檔產生「學生答案卷」與「解答卷」兩份 .docx（可選轉 PDF）。

兩份文件使用完全相同的幾何設定（固定欄寬、固定列高、固定行高），
解答卷的填空答案印在題號格、答案格留白，挖空答案格後即可疊在學生卷上比對。

用法：
    python3 make_answer_sheets.py config.json OUT_DIR [--pdf]
"""
import argparse
import json
import math
import os
import re
import shutil
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

# ---- 固定版面常數（mm）：學生卷與解答卷共用，改動會同時影響兩份 ----
TITLE_H = 12      # 標題行高
GAP_TITLE = 2     # 標題與基本資料列間距
ID_ROW_H = 12     # 班級/學號/姓名列高
GAP_SECTION = 3   # 區塊間距
HEADING_H = 7     # 大題標題行高
TAIL_H = 0.5      # 結尾極小段落
SAFETY = 8.5      # 保留給 Word/LibreOffice 排版差異的安全空間

DEFAULTS = {
    "paper": {"width_mm": 257, "height_mm": 364, "margin_lr_mm": 12, "margin_tb_mm": 10},
    "fonts": {"east_asia": "標楷體", "latin": "Times New Roman", "mono": "Courier New"},
    "key_title_suffix": "（解答）",
    "id_fields": ["班級", "學號", "姓名"],
    "id_widths_mm": None,
    "score_box": True,
    "key_color": "C00000",
    "shade": "E7E6E6",
}


# ---------------------------------------------------------------- helpers
class Ctx:
    def __init__(self, cfg):
        self.cfg = cfg
        self.ea = cfg["fonts"]["east_asia"]
        self.latin = cfg["fonts"]["latin"]
        self.mono = cfg["fonts"]["mono"]
        self.red = RGBColor.from_string(cfg["key_color"])
        self.gray = cfg["shade"]


def rich(ctx, p, s, size=12, bold=False, color=None, latin=None):
    """寫入文字；支援 _{下標} 與 ^{上標} 標記。"""
    for tok in re.split(r"(_\{[^}]*\}|\^\{[^}]*\})", s):
        if not tok:
            continue
        sub, sup = tok.startswith("_{"), tok.startswith("^{")
        r = p.add_run(tok[2:-1] if (sub or sup) else tok)
        r.font.size = Pt(size)
        r.bold = bold
        if sub:
            r.font.subscript = True      # 注意：不可再設定 superscript=None，否則會清掉下標
        elif sup:
            r.font.superscript = True
        if color is not None:
            r.font.color.rgb = color
        r.font.name = latin or ctx.latin
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), ctx.ea)


def mark_size(p, size):
    pPr = p._p.get_or_add_pPr()
    rpr = OxmlElement("w:rPr")
    for tag in ("w:sz", "w:szCs"):
        e = OxmlElement(tag)
        e.set(qn("w:val"), str(int(size * 2)))
        rpr.append(e)
    pPr.append(rpr)


def line(ctx, doc, text, mm, size=12, bold=False, align=None):
    """固定行高段落，確保兩份文件的垂直位置一致。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = Mm(mm)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.space_before = pf.space_after = Pt(0)
    mark_size(p, size)
    if align is not None:
        p.alignment = align
    if text:
        rich(ctx, p, text, size, bold)
    return p


def fix_table(t, widths):
    t.autofit = False
    t.alignment = 1
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", 30), ("left", 70), ("bottom", 0), ("right", 70)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "8")
        e.set(qn("w:color"), "000000")
        b.append(e)
    tblPr.append(b)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(round(w * 56.6929)))
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = Mm(w)


def set_h(row, h):
    row.height = Mm(h)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY


def shade(c, fill):
    tcPr = c._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), fill)
    tcPr.append(s)


def cell(ctx, c, text, size=12, bold=False, align="center", valign="center", fill=None):
    c.vertical_alignment = (WD_CELL_VERTICAL_ALIGNMENT.CENTER if valign == "center"
                            else WD_CELL_VERTICAL_ALIGNMENT.TOP)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == "center" else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    if text:
        rich(ctx, p, text, size, bold)
    if fill:
        shade(c, fill)
    return p


def key_block(ctx, c, items, size=11):
    """問答題解答：items 為 [kind, text]，kind = p 一般 / h 粗體小標 / m 等寬對齊"""
    for kind, s in items:
        p = c.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.space_before = Pt(3 if kind == "h" else 0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        if kind == "m":
            rich(ctx, p, s, size, False, ctx.red, latin=ctx.mono)
        else:
            rich(ctx, p, s, size, kind == "h", ctx.red)


def text_width(s):
    """粗估顯示寬度（中日韓字=1，其餘=0.5）"""
    s = re.sub(r"[_^]\{|\}", "", s)
    return sum(1 if ord(ch) > 0x2E80 else 0.5 for ch in s)


# ---------------------------------------------------------------- layout
def geometry(cfg):
    pa = cfg["paper"]
    W = pa["width_mm"] - 2 * pa["margin_lr_mm"]
    H = pa["height_mm"] - 2 * pa["margin_tb_mm"]
    fi = cfg["fill_in"]
    cols = fi.get("columns", 3)
    rows = fi.get("rows") or math.ceil(fi["count"] / cols)
    row_h = fi.get("row_height_mm", 9.5)
    fixed = (TITLE_H + GAP_TITLE + ID_ROW_H + GAP_SECTION + HEADING_H + rows * row_h
             + GAP_SECTION + HEADING_H + TAIL_H + SAFETY)
    written_h = H - fixed
    return W, H, cols, rows, row_h, written_h


def build(cfg, out, key):
    ctx = Ctx(cfg)
    W, H, NCOL, ROWS, ROW_H, WRITTEN_H = geometry(cfg)
    if WRITTEN_H < 60:
        sys.exit(f"問答區只剩 {WRITTEN_H:.0f} mm，請減少填空列數/列高或改用較大紙張")

    doc = Document()
    pa = cfg["paper"]
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(pa["width_mm"]), Mm(pa["height_mm"])
    sec.left_margin = sec.right_margin = Mm(pa["margin_lr_mm"])
    sec.top_margin = sec.bottom_margin = Mm(pa["margin_tb_mm"])
    sec.header_distance = sec.footer_distance = Mm(5)
    st = doc.styles["Normal"]
    st.font.name = ctx.latin
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), ctx.ea)
    st.paragraph_format.space_before = st.paragraph_format.space_after = Pt(0)

    # 標題
    title = cfg["title"] + (cfg["key_title_suffix"] if key else "")
    line(ctx, doc, title, TITLE_H, cfg.get("title_pt", 22), True, WD_ALIGN_PARAGRAPH.CENTER)
    line(ctx, doc, "", GAP_TITLE, 2)

    # 班級 / 學號 / 姓名 / 得分
    fields = cfg["id_fields"]
    label_w = 16
    score_w = 48 if cfg["score_box"] else 0
    vals = cfg["id_widths_mm"] or [(W - label_w * len(fields) - score_w) / len(fields)] * len(fields)
    widths = []
    for v in vals:
        widths += [label_w, v]
    if cfg["score_box"]:
        widths.append(W - sum(widths))
    t = doc.add_table(rows=1, cols=len(widths))
    fix_table(t, widths)
    set_h(t.rows[0], ID_ROW_H)
    cs = t.rows[0].cells
    for i, lab in enumerate(fields):
        cell(ctx, cs[2 * i], lab, 14, True, fill=ctx.gray)
        cell(ctx, cs[2 * i + 1], "")
    if cfg["score_box"]:
        cell(ctx, cs[-1], "得分：", 12, align="left")

    # 一、填空題
    fi = cfg["fill_in"]
    line(ctx, doc, "", GAP_SECTION, 3)
    line(ctx, doc, fi["key_heading"] if key else fi["heading"], HEADING_H, 13, True)
    num_w = fi.get("num_width_mm", 34)
    ans_w = W / NCOL - num_w
    t = doc.add_table(rows=ROWS, cols=NCOL * 2)
    fix_table(t, [num_w, ans_w] * NCOL)
    answers = {int(k): v for k, v in fi.get("answers", {}).items()}
    empty = []
    for i in range(ROWS):
        set_h(t.rows[i], ROW_H)
        cs = t.rows[i].cells
        for k in range(NCOL):
            n = k * ROWS + i + 1          # 直式編號：先排滿第一欄
            nc, ac = cs[2 * k], cs[2 * k + 1]
            if n > fi["count"]:
                empty.append((i, k))
                continue
            if key:
                p = cell(ctx, nc, "", align="left")
                rich(ctx, p, f"{n} ", 12, True)
                a = answers.get(n, "")
                # 長答案自動縮小，避免換行被固定列高切掉
                rich(ctx, p, " " + a, 11 if text_width(a) <= 6 else 9.5, True, ctx.red)
            else:
                cell(ctx, nc, str(n), 12, True, fill=ctx.gray)
            cell(ctx, ac, "")
    for i, k in empty:                    # 多出來的格子合併並塗灰
        r = t.rows[i].cells
        shade(r[2 * k].merge(r[2 * k + 1]), ctx.gray)

    # 二、問答/計算題：依 layout 設定把作答框排成多欄
    wr = cfg["written"]
    line(ctx, doc, "", GAP_SECTION, 3)
    line(ctx, doc, wr["key_heading"] if key else wr["heading"], HEADING_H, 13, True)
    qs = {q["id"]: q for q in wr["questions"]}
    lay = wr["layout"]
    col_w = lay["columns_mm"]
    col_w = [w * W / sum(col_w) for w in col_w]          # 依比例撐滿版心
    # 每欄的框依 weight 分配高度，再取所有邊界的聯集當作表格列
    bounds, boxes = {0.0, round(WRITTEN_H, 3)}, []
    for ci, col in enumerate(lay["columns"]):
        tot = sum(b["weight"] for b in col)
        y = 0.0
        for b in col:
            h = WRITTEN_H * b["weight"] / tot
            boxes.append((ci, y, y + h, qs[b["q"]]))
            y += h
            bounds.add(round(y, 3))
    edges = sorted(bounds)
    t = doc.add_table(rows=len(edges) - 1, cols=len(col_w))
    fix_table(t, col_w)
    for r in range(len(edges) - 1):
        set_h(t.rows[r], edges[r + 1] - edges[r])
    for ci, y0, y1, q in boxes:
        r0 = edges.index(round(y0, 3))
        r1 = edges.index(round(y1, 3)) - 1
        c = t.cell(r0, ci) if r0 == r1 else t.cell(r0, ci).merge(t.cell(r1, ci))
        cell(ctx, c, q["label"], 12, True, align="left", valign="top")
        if key:
            key_block(ctx, c, q.get("key", []))

    line(ctx, doc, "", TAIL_H, 1)   # 避免 Word 自動補段落把版面推到第二頁
    doc.save(out)


def to_pdf(path):
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        print("找不到 LibreOffice，略過 PDF 轉檔")
        return
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir",
                    os.path.dirname(path) or ".", path],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("config")
    ap.add_argument("outdir")
    ap.add_argument("--pdf", action="store_true", help="同時用 LibreOffice 轉出 PDF")
    a = ap.parse_args()
    with open(a.config, encoding="utf-8") as f:
        user = json.load(f)
    cfg = json.loads(json.dumps(DEFAULTS))
    for k, v in user.items():
        if isinstance(v, dict) and isinstance(cfg.get(k), dict):
            cfg[k].update(v)
        else:
            cfg[k] = v
    os.makedirs(a.outdir, exist_ok=True)
    base = cfg.get("output_name", cfg["title"])
    stu = os.path.join(a.outdir, base + ".docx")
    key = os.path.join(a.outdir, base + cfg.get("key_file_suffix", "_解答") + ".docx")
    build(cfg, stu, key=False)
    build(cfg, key, key=True)
    W, H, NCOL, ROWS, ROW_H, WH = geometry(cfg)
    print(f"學生卷：{stu}\n解答卷：{key}\n問答區高度 {WH:.0f} mm")
    if a.pdf:
        to_pdf(stu)
        to_pdf(key)


if __name__ == "__main__":
    main()
