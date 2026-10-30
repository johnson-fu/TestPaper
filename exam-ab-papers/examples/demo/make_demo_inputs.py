#!/usr/bin/env python3
"""產生範例用的座位表（seats.docx）、名冊（roster.xlsx）與附圖（img/graph.png）。姓名、學號皆為虛構。"""
import os
import docx, openpyxl
from PIL import Image, ImageDraw, ImageFont

os.makedirs("img", exist_ok=True)
# 座位表：2 列 × 4 行，A/B 交錯，一個空位
grid = [["王小明 (A)", "李小華 (B)", "陳大同 (A)", "X"],
        ["林美玲 (B)", "張志強 (A)", "黃雅婷 (B)", "吳建宏 (A)"]]
d = docx.Document()
t = d.add_table(rows=3, cols=5); t.style = "Table Grid"
t.cell(0, 0).text = "行別\n列號"; t.cell(0, 2).text = "(講桌)"
for r, row in enumerate(grid, 1):
    t.cell(r, 0).text = f"{r}."
    for c, v in enumerate(row, 1): t.cell(r, c).text = v
d.save("seats.docx")

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "名冊"
ws.append(["班級", "學號", "性別", "姓名"])
people = [("日電子三甲", "U0000001", "王小明"), ("日電子三甲", "U0000002", "李小華"), ("日電子三乙", "U0000003", "陳大同"),
          ("日電子三乙", "U0000004", "林美玲"), ("日資工二甲", "U0000005", "張志強"), ("日資工二甲", "U0000006", "黃雅婷"),
          ("日電機三甲", "U0000007", "吳建宏")]
for cls, sid, name in people: ws.append([cls, sid, "", name])
wb.save("roster.xlsx")

# 附圖：正方形加一條對角線（5 條邊）
im = Image.new("RGB", (600, 600), "white"); g = ImageDraw.Draw(im)
P = {"a": (120, 120), "b": (480, 120), "c": (480, 480), "d": (120, 480)}
for u, v in [("a", "b"), ("b", "c"), ("c", "d"), ("d", "a"), ("a", "c")]:
    g.line([P[u], P[v]], fill="black", width=6)
try: font = ImageFont.truetype("DejaVuSans.ttf", 48)
except OSError: font = ImageFont.load_default()
for k, (x, y) in P.items():
    g.ellipse([x - 16, y - 16, x + 16, y + 16], fill="black")
    g.text((x + (-60 if x < 300 else 30), y + (-70 if y < 300 else 20)), k, fill="black", font=font)
im.save("img/graph.png")
print("seats.docx、roster.xlsx、img/graph.png 已產生")
