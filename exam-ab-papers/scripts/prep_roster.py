#!/usr/bin/env python3
"""座位表（.docx）＋ 名冊（.xlsx）→ students.json

座位表：Word 表格，第一列為行別（講桌）、第一欄為列號；其餘每格為「姓名 (A)」或「姓名（B）」，空位寫 X 或留白。
名冊：第一列為標題，需含「班級」「學號」「姓名」三欄（欄位順序不拘，可有其他欄）。

排序（--order）：
  column（預設）＝由上而下、由左至右：先排第 1 行的第 1～N 列，再換第 2 行……
  row            ＝由左至右、由上而下：先排第 1 列的第 1～N 行，再換第 2 列……

用法：python3 prep_roster.py 座位表.docx 名冊.xlsx [--order column|row] [-o students.json]
"""
import argparse, json, re, sys
import docx, openpyxl

ap = argparse.ArgumentParser()
ap.add_argument("seat_docx"); ap.add_argument("roster_xlsx")
ap.add_argument("--order", choices=["column", "row"], default="column")
ap.add_argument("--sheet", help="名冊工作表名稱（預設第一張）")
ap.add_argument("-o", "--out", default="students.json")
a = ap.parse_args()

# ---------- 座位表 ----------
d = docx.Document(a.seat_docx)
if not d.tables: sys.exit("座位表找不到表格")
t = d.tables[0]
seats = []
pat = re.compile(r"^(.+?)\s*[（(]\s*([ABab])\s*[)）]$")
for r, row in enumerate(t.rows[1:], 1):
    cells = row.cells[1:]
    for c, cell in enumerate(cells, 1):
        txt = re.sub(r"\s+", " ", cell.text).strip()
        if not txt or txt.upper() in ("X", "×", "空"): continue
        m = pat.match(txt)
        if not m: sys.exit(f"座位表第 {r} 列第 {c} 行無法辨識：{txt!r}（應為「姓名 (A)」）")
        seats.append({"row": r, "col": c, "name": m.group(1).strip(), "ver": m.group(2).upper()})

# ---------- 名冊 ----------
wb = openpyxl.load_workbook(a.roster_xlsx, data_only=True)
ws = wb[a.sheet] if a.sheet else wb.worksheets[0]
header = [str(v).strip() if v is not None else "" for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
try:
    ic, ii, iname = header.index("班級"), header.index("學號"), header.index("姓名")
except ValueError:
    sys.exit(f"名冊第一列需含「班級」「學號」「姓名」，目前為：{header}")
roster = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    name = row[iname]
    if not name: continue
    roster.setdefault(str(name).strip(), []).append({"cls": str(row[ic] or "").strip(), "id": str(row[ii] or "").strip()})

# ---------- 比對 ----------
problems = []
dups = [n for n, v in roster.items() if len(v) > 1]
if dups: problems.append(f"名冊同名：{dups}（請在座位表姓名後加註以區分，並修改本程式比對規則）")
seat_names = [s["name"] for s in seats]
if len(set(seat_names)) != len(seat_names): problems.append("座位表有重複姓名")
missing = [n for n in seat_names if n not in roster]
if missing: problems.append(f"座位表有、名冊沒有：{missing}")
absent = [n for n in roster if n not in seat_names]
if absent: print(f"提醒：名冊有、座位表沒有（不產生考卷）：{absent}")
if problems: sys.exit("\n".join(problems))

for s in seats: s.update(roster[s["name"]][0])
seats.sort(key=(lambda s: (s["col"], s["row"])) if a.order == "column" else (lambda s: (s["row"], s["col"])))
json.dump(seats, open(a.out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
na = sum(s["ver"] == "A" for s in seats)
print(f"{len(seats)} 人（A 卷 {na}、B 卷 {len(seats) - na}），順序：{a.order} → {a.out}")
print(" → ".join(f"{s['name']}({s['ver']})" for s in seats))
